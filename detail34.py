from pptx import Presentation
from pptx.util import Emu, Pt
prs = Presentation(r"C:\Users\shaik\Downloads\THREX_SIH26145_work.pptx")
out = open("slide34_detail.txt","w",encoding="utf-8")
def dump(shapes, out, ind=0):
    for sh in shapes:
        geo = f"({Emu(sh.left).inches:6.2f},{Emu(sh.top).inches:5.2f} {Emu(sh.width).inches:5.2f}x{Emu(sh.height).inches:5.2f})"
        out.write(" "*ind + f"--- id={sh.shape_id} name={sh.name} type={sh.shape_type} geo={geo}\n")
        if sh.shape_type == 6:
            dump(sh.shapes, out, ind+3); continue
        if sh.has_text_frame:
            tf = sh.text_frame
            out.write(" "*ind + f"    wrap={tf.word_wrap} autosize={tf.auto_size} margins L{Emu(tf.margin_left).inches:.2f} R{Emu(tf.margin_right).inches:.2f} T{Emu(tf.margin_top).inches:.2f} B{Emu(tf.margin_bottom).inches:.2f}\n")
            for pi, p in enumerate(tf.paragraphs):
                out.write(" "*ind + f"    P{pi} align={p.alignment} lvl={p.level} sz={p.font.size and p.font.size.pt} spc_b={p.space_before} spc_a={p.space_after} ls={p.line_spacing} txt={p.text!r}\n")
                for ri, r in enumerate(p.runs):
                    c = None
                    try: c = r.font.color.rgb
                    except Exception: pass
                    out.write(" "*ind + f"       R{ri} sz={r.font.size and r.font.size.pt} b={r.font.bold} i={r.font.italic} name={r.font.name} color={c} txt={r.text!r}\n")
for idx in (2,3):
    out.write("\n"+"#"*60 + f" SLIDE {idx+1} " + "#"*60 + "\n")
    dump(prs.slides[idx].shapes, out)
out.close()
print("ok")
