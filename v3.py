from pptx import Presentation
prs = Presentation(r"C:\Users\shaik\Downloads\THREX_SIH2026_v3.pptx")
out=open("v3_text.txt","w",encoding="utf-8")
def walk(shapes,out,ind=1):
    for sh in shapes:
        if sh.shape_type == 6:
            walk(sh.shapes,out,ind+2); continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.write(" "*ind+f"{sh.shape_id}: "+sh.text_frame.text.replace(chr(10)," || ")+"\n")
for i,s in enumerate(prs.slides,1):
    out.write(f"\n##### SLIDE {i} #####\n"); walk(s.shapes,out)
out.close()
import re
txt=open("v3_text.txt",encoding="utf-8").read()
for k in ["1,978.6","162.7","PS \u00b7 NTRO","RISK","Risk","Feasibility","Not built yet","Weaknesses","13,394","2,800","Three correctness"]:
    print(k, "->", [i+1 for i,l in enumerate(txt.splitlines()) if k in l][:6])
