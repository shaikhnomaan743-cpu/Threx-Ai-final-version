from pptx import Presentation
from pptx.util import Emu
files = [
 r"C:\Users\shaik\Downloads\THREX_SIH26145_work.pptx",
 r"C:\Users\shaik\Downloads\THREX_SIH26145_updated.pptx",
 r"C:\Users\shaik\Downloads\THREX_SIH2026_v3.pptx",
 r"C:\Users\shaik\Downloads\THREX_SIH2026_v3 (1).pptx",
]
out = open("deck_text.txt","w",encoding="utf-8")
def walk(shapes, out, ind=1):
    for sh in shapes:
        if sh.shape_type == 6:
            out.write(" "*ind + f"[group {sh.name}]\n")
            walk(sh.shapes, out, ind+2)
        else:
            t = ""
            if sh.has_text_frame:
                t = " | ".join(pp.text.replace(chr(10)," / ") for pp in sh.text_frame.paragraphs if pp.text.strip())
            out.write(" "*ind + f"{sh.name}: {t}\n" if t else " "*ind + f"{sh.name}: <{sh.shape_type}>\n")
for f in files:
    prs = Presentation(f)
    out.write("="*100+"\n")
    out.write(f"FILE: {f}  slides={len(prs.slides)}  {Emu(prs.slide_width).inches:.2f}x{Emu(prs.slide_height).inches:.2f}in\n")
    out.write("="*100+"\n")
    for i, s in enumerate(prs.slides, 1):
        out.write(f"\n##### SLIDE {i} #####\n")
        walk(s.shapes, out)
out.close()
print("ok")
