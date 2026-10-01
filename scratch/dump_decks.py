from pptx import Presentation
from pptx.util import Emu
files = [
 r"C:\Users\shaik\Downloads\THREX_SIH26145_work.pptx",
 r"C:\Users\shaik\Downloads\THREX_SIH26145_updated.pptx",
]
out = open("deck_dump.txt","w",encoding="utf-8")
for f in files:
    prs = Presentation(f)
    out.write("="*110+"\n")
    out.write(f"FILE: {f} | slides: {len(prs.slides)} | size: {prs.slide_width} x {prs.slide_height}\n")
    out.write("="*110+"\n")
    for i, s in enumerate(prs.slides, 1):
        out.write(f"--- SLIDE {i} ---\n")
        for sh in s.shapes:
            t = ""
            if sh.has_text_frame:
                t = " || ".join(pp.text for pp in sh.text_frame.paragraphs)
            else:
                t = f"<{sh.shape_type}>"
            out.write(f"   [{sh.shape_id}] {sh.name} ({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f} {Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f}) :: {t[:900]}\n")
        out.write("\n")
out.close()
print("done")
