from pptx import Presentation
from pptx.util import Emu
prs = Presentation(r"C:\Users\shaik\Downloads\THREX_SIH26145_work.pptx")
out = open("focus.txt","w",encoding="utf-8")
targets = {1:[63], 2:[7,10,15], 3:[23,25,29,31,36,37,38,42]}
for si, ids in targets.items():
    s = prs.slides[si]
    out.write("\n"+"="*40+f" SLIDE {si+1} "+"="*40+"\n")
    for sh in s.shapes:
        if sh.shape_id not in ids: continue
        geo = f"({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f} {Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f})"
        out.write(f"id={sh.shape_id} name={sh.name} type={sh.shape_type} geo={geo}\n")
        if not sh.has_text_frame: continue
        tf = sh.text_frame
        out.write(f"  wrap={tf.word_wrap} autosize={tf.auto_size} margins L{Emu(tf.margin_left).inches:.2f} R{Emu(tf.margin_right).inches:.2f} T{Emu(tf.margin_top).inches:.2f} B{Emu(tf.margin_bottom).inches:.2f}\n")
        for pi,p in enumerate(tf.paragraphs):
            out.write(f"  P{pi} lvl={p.level} align={p.alignment} sz={p.font.size and p.font.size.pt} spcB={p.space_before} spcA={p.space_after} ls={p.line_spacing}\n")
            for ri,r in enumerate(p.runs):
                c=None
                try: c=r.font.color.rgb
                except Exception: pass
                out.write(f"     R{ri} sz={r.font.size and r.font.size.pt} b={r.font.bold} name={r.font.name} color={c} len={len(r.text)} txt={r.text[:120]!r}\n")
out.close(); print('ok')
