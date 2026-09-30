from pptx import Presentation
import glob, os
phrases = ["Feasibility: it already runs","Feasibility and Viability","Risk 2","precision is 0.30","13,394","2,800","1,978.6","162.7","270,000","5,300","JA4","QUIC","synthetic","NetFlow","Caffeinated Coders","PS \u00b7 NTRO","SIH26145 \u00b7 NTRO","Research and References","Technical Approach","unverified","static","roadmap"]
def alltext(prs):
    per = []
    for s in prs.slides:
        buf = []
        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == 6:
                    walk(sh.shapes)
                else:
                    if sh.has_text_frame: buf.append(sh.text_frame.text)
                    if sh.has_table:
                        for r in sh.table.rows:
                            buf.append(" ".join(c.text for c in r.cells))
        walk(s.shapes)
        per.append("\n".join(buf))
    return per
files = sorted(glob.glob(r"C:\Users\shaik\Downloads\*.pptx"))
out = open("phrase_matrix.txt","w",encoding="utf-8")
for f in files:
    try: prs = Presentation(f)
    except Exception as e:
        out.write(f"ERR {os.path.basename(f)}: {e}\n"); continue
    slides = alltext(prs)
    joined = "\n".join(slides)
    hits = {}
    for p in phrases:
        idx = [i+1 for i,t in enumerate(slides) if p.lower() in t.lower()]
        if idx: hits[p] = idx
    out.write(f"\n### {os.path.basename(f)} | slides={len(slides)} | bytes={os.path.getsize(f)}\n")
    for p, ix in hits.items():
        out.write(f"     {p!r:34s} -> slides {ix}\n")
out.close()
print("ok")
