from pptx import Presentation
from pptx.util import Emu
prs = Presentation(r"C:\Users\shaik\Downloads\THREX_SIH26145_work.pptx")
out=open("focus2.txt","w",encoding="utf-8")
for sh in prs.slides[3].shapes:
    if sh.shape_id not in (29,31,35,36,39): continue
    geo=f"({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f} {Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f})"
    out.write(f"\nid={sh.shape_id} name={sh.name} type={sh.shape_type} geo={geo}\n")
    if not sh.has_text_frame: continue
    tf=sh.text_frame
    out.write(f"  wrap={tf.word_wrap} autosize={tf.auto_size} margins L{Emu(tf.margin_left).inches:.2f} T{Emu(tf.margin_top).inches:.2f} B{Emu(tf.margin_bottom).inches:.2f}\n")
    for pi,p in enumerate(tf.paragraphs):
        out.write(f"  P{pi} lvl={p.level} align={p.alignment} spcB={p.space_before} spcA={p.space_after} ls={p.line_spacing}\n")
        for ri,r in enumerate(p.runs):
            c=None
            try: c=r.font.color.rgb
            except Exception: pass
            out.write(f"     R{ri} sz={r.font.size and r.font.size.pt} b={r.font.bold} name={r.font.name} color={c} len={len(r.text)}\n        {r.text!r}\n")
out.close(); print('ok')
