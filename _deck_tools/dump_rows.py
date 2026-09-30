"""Dump group internals + constraints rows on slide 4 (encoding-safe)."""
import sys, os
from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from wrapmeasure import wrapped_lines

EMU_PT = 12700.0
DECK = r"C:\Users\shaik\Downloads\THREX_SIH26145_work.pptx"
prs = Presentation(DECK)
slide = prs.slides[3]
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "rows_s4.txt")

with open(OUT, "w", encoding="utf-8") as f:
    def w(s=""):
        f.write(s + "\n")

    # groups: internals
    for sh in slide.shapes:
        if sh.shape_type != 6 or sh.shape_id not in (26, 32):
            continue
        w("=" * 90)
        w(f"GROUP id={sh.shape_id} top={sh.top/EMU_PT:.2f} left={sh.left/EMU_PT:.2f} "
          f"w={sh.width/EMU_PT:.2f} h={sh.height/EMU_PT:.2f} chOff={sh._element.chOff} chExt={sh._element.chExt}")
        for c in sh.shapes:
            w(f"   child id={c.shape_id} type={c.shape_type} top={c.top/EMU_PT:.2f} "
              f"left={c.left/EMU_PT:.2f} w={c.width/EMU_PT:.2f} h={c.height/EMU_PT:.2f}")
            try:
                prst = c._element.spPr.prstGeom.get("prst")
                w(f"        prstGeom={prst}")
            except Exception as e:
                w(f"        prstGeom=<none> {e}")
            try:
                ln = c._element.spPr.find(
                    '{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
                if ln is not None:
                    wl = ln.find('{http://schemas.openxmlformats.org/drawingml/2006/main}w')
                    w(f"        line w={wl.get('emu') if wl is not None else None}")
            except Exception:
                pass

    # constraint rows
    for sh in slide.shapes:
        if sh.shape_id not in (35, 36, 37, 38, 40, 41):
            continue
        w("=" * 90)
        w(f"id={sh.shape_id} name={sh.name} top={sh.top/EMU_PT:.2f} left={sh.left/EMU_PT:.2f} "
          f"w={sh.width/EMU_PT:.2f} h={sh.height/EMU_PT:.2f} bottom={(sh.top+sh.height)/EMU_PT:.2f}")
        if not sh.has_text_frame:
            continue
        tf = sh.text_frame
        avail = (sh.width - tf.margin_left - tf.margin_right) / EMU_PT
        tot = 0.0
        for pi, p in enumerate(tf.paragraphs):
            if not p.runs:
                w(f"   P{pi} <empty>")
                continue
            sz = p.runs[0].font.size
            sz = sz.pt if sz else 18.0
            n = len(wrapped_lines(p.text, sz, avail))
            spcb = (p.space_before / EMU_PT) if p.space_before is not None else 0.0
            tot += spcb + n * 1.22 * sz
            w(f"   P{pi} sz={sz} spcB={spcb:.2f} lines={n} chars={len(p.text)}")
            w(f"        {p.text}")
        w(f"   --> predicted height {tot:.2f}pt, text bottom y={sh.top/EMU_PT + tot:.2f}")

print("wrote", OUT)
