"""Slide-4 right-column analysis: exact geometry + rendered line counts."""
import sys, os
from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from wrapmeasure import wrapped_lines, text_width

EMU_PT = 12700.0  # EMU per point
LH_RATIO = 1.22   # Calibri single-spacing line-height ratio (validated against COM)

DECK = r"C:\Users\shaik\Downloads\THREX_SIH26145_work.pptx"
WANT = [26, 29, 31, 32, 35, 36, 37, 38, 39, 40, 41, 42, 20, 23, 25, 43]

prs = Presentation(DECK)
slide = prs.slides[3]


def emu2pt(v):
    return v / EMU_PT


for sh in slide.shapes:
    if sh.shape_id not in WANT:
        continue
    print("=" * 100)
    print(f"id={sh.shape_id} name={sh.name} type={sh.shape_type} "
          f"top={emu2pt(sh.top):.2f} left={emu2pt(sh.left):.2f} "
          f"w={emu2pt(sh.width):.2f} h={emu2pt(sh.height):.2f} "
          f"bottom={emu2pt(sh.top + sh.height):.2f}")
    if not sh.has_text_frame:
        continue
    tf = sh.text_frame
    width_pt = emu2pt(sh.width) - emu2pt(tf.margin_left) - emu2pt(tf.margin_right)
    print(f"  text width available = {width_pt:.2f}pt  autosize={tf.auto_size} "
          f"wrap={tf.word_wrap} spcB_total/pt={0}")
    total = 0.0
    for pi, p in enumerate(tf.paragraphs):
        if p.runs:
            sz = p.runs[0].font.size
            sz = sz.pt if sz is not None else 18.0
            txt = p.text
            n = len(wrapped_lines(txt, sz, width_pt))
            spcb = emu2pt(p.space_before) if p.space_before is not None else 0.0
            lh = LH_RATIO * sz
            total += spcb + n * lh
            print(f"   P{pi} sz={sz} spcB={spcb:.2f}pt lines={n} lh={lh:.2f} "
                  f"subtotal={spcb + n * lh:.2f}  chars={len(txt)}")
            print(f"        {txt}")
        else:
            print(f"   P{pi} <empty>")
    print(f"  PREDICTED TEXT HEIGHT = {total:.2f}pt "
          f"(y {emu2pt(sh.top):.2f} -> {emu2pt(sh.top) + total:.2f})")
    if tf.auto_size is None:
        pass
